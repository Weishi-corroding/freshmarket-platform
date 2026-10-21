"""Build the enterprise Word deliverables (SRS / SDD / test report) from docs/*.md.

Formatting: A4, all text black; body 宋体 五号 + Times New Roman, headings 黑体, code Consolas; tables use plain black
borders without shading. Only images and diagrams carry colour.
Mermaid blocks are replaced by PNGs listed in docs/diagrams/mermaid-manifest.json, keyed by the first 16 hex chars of
SHA-256 over the stripped block text (render missing ones with scripts/render-mermaid.mjs).
Usage: python scripts/generate-enterprise-docs.py [requirements design test-report]
然后运行 scripts/render-documents.ps1 由 Word 更新目录与页码并导出 PDF。
"""
from pathlib import Path
import hashlib,json,re,sys
from docx import Document
from docx.enum.section import WD_ORIENT,WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt,RGBColor
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/'docs';OUT=ROOT/'deliverables';OUT.mkdir(exist_ok=True)
DIAGRAMS=DOCS/'diagrams'
MANIFEST=json.loads((DIAGRAMS/'mermaid-manifest.json').read_text('utf-8'))
BLACK='000000'
SONG,HEI,LATIN,MONO='宋体','黑体','Times New Roman','Consolas'
SPECS={
    'requirements':('鲜境_需求文档','软件需求规格说明书','Software Requirements Specification','FM-SRS-001'),
    'design':('鲜境_设计文档','系统设计说明书','Software Design Description','FM-SDD-001'),
    'test-report':('鲜境_测试报告','系统测试报告','Software Test Report','FM-STR-001'),
}
META={'版本':'3.0','日期':'2026-10-21','状态':'正式稿','适用版本':'main 分支（2026-10-21）'}

def mermaid_key(block):return hashlib.sha256('\n'.join(block).strip().encode('utf-8')).hexdigest()[:16]

def set_fonts(rpr,latin,cjk):
    rf=rpr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rpr.append(rf)
    for k in list(rf.attrib):
        if k.endswith('Theme') or k.endswith('theme'):del rf.attrib[k]  # 主题字体优先级高于显式字体，必须移除
    rf.set(qn('w:ascii'),latin);rf.set(qn('w:hAnsi'),latin);rf.set(qn('w:cs'),latin);rf.set(qn('w:eastAsia'),cjk)

def fonts(run,size=None,bold=None,mono=False,cjk=SONG,italic=None):
    run.font.name=MONO if mono else LATIN
    set_fonts(run._element.get_or_add_rPr(),MONO if mono else LATIN,cjk)
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    if italic is not None:run.italic=italic
    run.font.color.rgb=RGBColor.from_string(BLACK)
    return run

def border(el,sides,sz='4',space=None):
    tag=qn('w:tblBorders') if el.tag==qn('w:tblPr') else qn('w:pBdr')
    b=OxmlElement('w:tblBorders' if tag==qn('w:tblBorders') else 'w:pBdr')
    for side in sides:
        e=OxmlElement(f'w:{side}');e.set(qn('w:val'),'single');e.set(qn('w:sz'),sz);e.set(qn('w:color'),BLACK)
        if space:e.set(qn('w:space'),space)
        b.append(e)
    el.append(b)

def field(paragraph,instr,placeholder='',size=9):
    run=paragraph.add_run();b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');b.set(qn('w:dirty'),'true');run._r.append(b)
    run=paragraph.add_run();t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=instr;run._r.append(t)
    run=paragraph.add_run();s=OxmlElement('w:fldChar');s.set(qn('w:fldCharType'),'separate');run._r.append(s)
    fonts(paragraph.add_run(placeholder),size)
    run=paragraph.add_run();e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');run._r.append(e)

INLINE=re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')
def inline(paragraph,text,size=None,bold=False):
    text=text.replace('<br>','\n').replace('<br/>','\n')
    for part in INLINE.split(text):
        if not part:continue
        if part.startswith('**') and part.endswith('**'):fonts(paragraph.add_run(part[2:-2]),size,True)
        elif part.startswith('`') and part.endswith('`'):fonts(paragraph.add_run(part[1:-1]),(size or 10.5)-1,bold,mono=True)
        elif part.startswith('['):
            m=re.match(r'\[([^\]]+)\]\(([^)]+)\)',part);fonts(paragraph.add_run(m[1]),size,bold)
            if not m[2].startswith('#'):fonts(paragraph.add_run(f'（{m[2]}）'),size,False)
        else:fonts(paragraph.add_run(part),size,bold)

def cell_margins(table):
    mar=OxmlElement('w:tblCellMar')
    for side,v in (('top',40),('bottom',40),('left',80),('right',80)):
        el=OxmlElement(f'w:{side}');el.set(qn('w:w'),str(v));el.set(qn('w:type'),'dxa');mar.append(el)
    table._tbl.tblPr.append(mar)

class Builder:
    def __init__(self,stem):
        self.stem=stem;self.file,self.title,self.en,self.code=SPECS[stem]
        self.doc=Document();self.figure=0;self.last_heading=''
        self.setup();self.cover();self.toc()
    def setup(self):
        d=self.doc;sec=d.sections[0];sec.page_height,sec.page_width=Cm(29.7),Cm(21)
        sec.top_margin=sec.bottom_margin=Cm(2.5);sec.left_margin=sec.right_margin=Cm(2.5)
        st=d.styles['Normal'];st.font.name=LATIN;st.font.size=Pt(10.5);st.font.color.rgb=RGBColor.from_string(BLACK)
        set_fonts(st.element.get_or_add_rPr(),LATIN,SONG)
        st.paragraph_format.line_spacing=1.3;st.paragraph_format.space_after=Pt(3)
        for name,size,before,after in (('Heading 1',16,18,10),('Heading 2',14,12,6),('Heading 3',12,8,4)):
            h=d.styles[name];h.font.name=LATIN;h.font.size=Pt(size);h.font.bold=True;h.font.italic=False
            h.font.color.rgb=RGBColor.from_string(BLACK);set_fonts(h.element.get_or_add_rPr(),LATIN,HEI)
            h.paragraph_format.space_before=Pt(before);h.paragraph_format.space_after=Pt(after);h.paragraph_format.keep_with_next=True
            h.paragraph_format.line_spacing=1.2
        d.styles['List Bullet'].font.color.rgb=RGBColor.from_string(BLACK)
        sec.different_first_page_header_footer=True
        self.header_footer(sec)
    def header_footer(self,sec):
        sec.header.is_linked_to_previous=sec.footer.is_linked_to_previous=False
        hp=sec.header.paragraphs[0];hp.text='';hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        fonts(hp.add_run(f'鲜境 FreshMarket {self.title}    {self.code}  V{META["版本"]}'),9)
        border(hp._p.get_or_add_pPr(),('bottom',),'4','1')
        fp=sec.footer.paragraphs[0];fp.text='';fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        fonts(fp.add_run('第 '),9);field(fp,'PAGE','1');fonts(fp.add_run(' 页  共 '),9);field(fp,'NUMPAGES','1');fonts(fp.add_run(' 页'),9)
    def centered(self,text,size,bold=False,cjk=HEI,before=0,after=6):
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after)
        fonts(p.add_run(text),size,bold,cjk=cjk);return p
    def cover(self):
        self.centered('鲜境 FreshMarket 生鲜电商平台',16,True,before=110)
        self.centered(self.title,26,True,before=24)
        self.centered(self.en,14,False,cjk=SONG,after=90)
        t=self.doc.add_table(rows=0,cols=2);t.alignment=WD_TABLE_ALIGNMENT.CENTER;border(t._tbl.tblPr,('top','left','bottom','right','insideH','insideV'));cell_margins(t)
        for k,v in [('文档编号',self.code),('文档版本',META['版本']),('编制日期',META['日期']),('文档状态',META['状态']),('适用版本',META['适用版本']),('密    级','内部')]:
            c=t.add_row().cells;c[0].width=Cm(3.5);c[1].width=Cm(7.5)
            for cell,text,bold in ((c[0],k,True),(c[1],v,False)):
                p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0);fonts(p.add_run(text),10.5,bold)
        self.centered('鲜境 FreshMarket 项目组',12,False,cjk=SONG,before=80)
        p=self.centered('本文档依据当前源代码、脚本与 2026-10-21 实际执行结果编制。',9,False,cjk=SONG)
        p.add_run().add_break(WD_BREAK.PAGE)
    def toc(self):
        self.centered('目    录',16,True,after=12)
        field(self.doc.add_paragraph(),'TOC \\o "1-3" \\h \\z \\u','（打开文档后右键“更新域”生成目录）',10.5)
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    def caption(self,text):
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(10)
        self.figure+=1;fonts(p.add_run(f'图 {self.figure}  {text}'),9)
    def section(self,landscape):
        sec=self.doc.add_section(WD_SECTION.NEW_PAGE)
        w,h=Cm(29.7),Cm(21)
        sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        sec.page_width,sec.page_height=(w,h) if landscape else (h,w)
        sec.different_first_page_header_footer=False;self.header_footer(sec);return sec
    def image(self,path,caption):
        w,h=Image.open(path).size;wide=w/h>=1.45 and w>=3000
        if wide:self.section(True)
        sec=self.doc.sections[-1];maxw=(sec.page_width-sec.left_margin-sec.right_margin)/914400*2.54
        maxh=(sec.page_height-sec.top_margin-sec.bottom_margin)/914400*2.54-2.2
        css_w=w/2/96*2.54  # 按 2 倍设备像素渲染 → 自然尺寸（cm）
        width=min(maxw,css_w,maxh*w/h)
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(path),width=Cm(width))
        self.caption(caption)
        if wide:self.section(False)
    def table(self,rows):
        cols=max(len(r) for r in rows);t=self.doc.add_table(rows=0,cols=cols);t.alignment=WD_TABLE_ALIGNMENT.CENTER
        border(t._tbl.tblPr,('top','left','bottom','right','insideH','insideV'));cell_margins(t);size=8.5 if cols>=6 else 9
        for i,row in enumerate(rows):
            cells=t.add_row().cells
            for j in range(cols):
                p=cells[j].paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1
                if i==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                inline(p,row[j] if j<len(row) else '',size,bold=(i==0))
            trPr=t.rows[i]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:tblHeader' if i==0 else 'w:cantSplit'))
        self.doc.add_paragraph().paragraph_format.space_after=Pt(2)
    def code_block(self,lines):
        for k,line in enumerate(lines):
            p=self.doc.add_paragraph();pf=p.paragraph_format
            pf.space_after=Pt(0) if k<len(lines)-1 else Pt(8);pf.space_before=Pt(0);pf.line_spacing=1.0;pf.left_indent=Cm(.4);pf.right_indent=Cm(.2)
            border(p._p.get_or_add_pPr(),('top','left','bottom','right'),'4','4')  # 连续段落的相同边框合并为一个框
            fonts(p.add_run(line or ' '),8.5,mono=True)
    def numbered(self,n,text):
        p=self.doc.add_paragraph();pf=p.paragraph_format;pf.left_indent=Cm(.75);pf.first_line_indent=Cm(-.75)
        fonts(p.add_run(f'{n}. '),10.5);inline(p,text)  # 直接写序号，避免 List Number 在全文连续编号
    def render(self,text):
        lines=text.splitlines();i=0
        while i<len(lines):
            raw=lines[i];s=raw.strip();i+=1
            if not s:continue
            if s.startswith('```'):
                lang=s[3:].strip();block=[]
                while i<len(lines) and not lines[i].strip().startswith('```'):block.append(lines[i]);i+=1
                i+=1
                if lang=='mermaid':
                    key=mermaid_key(block)
                    if key not in MANIFEST:raise SystemExit(f'{self.stem}: mermaid block {key} has no rendered PNG; run node scripts/render-mermaid.mjs first')
                    self.image(DIAGRAMS/MANIFEST[key],self.last_heading)
                else:self.code_block(block)
                continue
            if s.startswith('|'):
                rows=[]
                while True:
                    cells=[c.strip() for c in re.split(r'(?<!\\)\|',s.strip().strip('|'))]
                    cells=[c.replace('\\|','|') for c in cells]
                    if not all(re.fullmatch(r':?-{2,}:?',c) for c in cells if c) or not any(cells):rows.append(cells)
                    if i>=len(lines) or not lines[i].strip().startswith('|'):break
                    s=lines[i].strip();i+=1
                self.table(rows);continue
            m=re.match(r'!\[([^\]]*)\]\(([^)]+)\)',s)
            if m:self.image((DOCS/m[2]).resolve(),m[1] or self.last_heading);continue
            m=re.match(r'(#{1,4})\s+(.+)',s)
            if m:
                level=len(m[1])
                if level==1:continue
                title=m[2].replace('`','').replace('**','')
                self.last_heading=re.sub(r'^(附录\s*\S+|[\d.]+)\s*','',title)
                h=self.doc.add_heading('',min(level-1,3));fonts(h.add_run(title),None,True,cjk=HEI);continue
            if re.match(r'[-*] ',s):p=self.doc.add_paragraph(style='List Bullet');inline(p,s[2:]);continue
            m=re.match(r'(\d+)\. (.*)',s)
            if m:self.numbered(m[1],m[2]);continue
            if s.startswith('>'):
                p=self.doc.add_paragraph();p.paragraph_format.left_indent=Cm(.75);border(p._p.get_or_add_pPr(),('left',),'8','8');inline(p,s.lstrip('> '));continue
            para=[s]
            while i<len(lines) and lines[i].strip() and not re.match(r'(```|\||#|[-*] |\d+\. |>|!\[)',lines[i].strip()):para.append(lines[i].strip());i+=1
            p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(21);inline(p,''.join(para))
    def save(self):
        cp=self.doc.core_properties;cp.title=f'鲜境 FreshMarket {self.title}';cp.subject=self.code;cp.author='鲜境 FreshMarket 项目组';cp.keywords='FreshMarket'
        path=OUT/f'{self.file}.docx';self.doc.save(path);return path

if __name__=='__main__':
    for stem in (sys.argv[1:] or list(SPECS)):
        b=Builder(stem);b.render((DOCS/f'{stem}.md').read_text('utf-8'));print(b.save(),f'figures={b.figure}')
