from pathlib import Path
import sys,re,json,subprocess,zipfile
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.tools'))
from PIL import Image,ImageOps,ImageDraw,ImageFont
from docx import Document
from pptx import Presentation
OUT=ROOT/'deliverables'
report={'documents':[],'presentation':{},'validation':'OOXML integrity, real Office PDF page counts, PowerPoint text bounds and render previews'}
for file in sorted(OUT.glob('*.docx')):
    with zipfile.ZipFile(file) as archive:assert archive.testzip() is None
    doc=Document(file);text='\n'.join(p.text for p in doc.paragraphs)
    assert len(text)>300 and '鲜境' in text
    pdf=file.with_suffix('.pdf');assert pdf.exists() and pdf.stat().st_size>10000
    info=subprocess.run(['pdfinfo',str(pdf)],capture_output=True,text=True,encoding='utf-8',errors='replace',check=True).stdout
    pages=int(re.search(r'Pages:\s+(\d+)',info)[1]);report['documents'].append({'file':file.name,'bytes':file.stat().st_size,'pdfPages':pages,'tables':len(doc.tables),'images':len(doc.inline_shapes)})
deck=next(OUT.glob('*.pptx'));prs=Presentation(deck);overflow=[]
for i,s in enumerate(prs.slides,1):
    for shape in s.shapes:
        if shape.left<0 or shape.top<0 or shape.left+shape.width>prs.slide_width+10 or shape.top+shape.height>prs.slide_height+10:overflow.append({'slide':i,'shape':shape.name})
assert not overflow,overflow
report['presentation']={'file':deck.name,'slides':len(prs.slides),'outsideCanvas':overflow,'actualOfficeExport':deck.with_suffix('.pdf').exists()}
slides=sorted((ROOT/'.runtime/ppt-preview').glob('*.PNG'),key=lambda p:int(re.search(r'(\d+)',p.name)[1]));assert len(slides)==len(prs.slides)
canvas=Image.new('RGB',(1600,280*((len(slides)+2)//3)),'#e9f0f0');draw=ImageDraw.Draw(canvas);font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',17)
for i,p in enumerate(slides):
    im=Image.open(p);im.thumbnail((510,247));x=(i%3)*533+10;y=(i//3)*280+8;canvas.paste(im,(x,y));draw.text((x,y+249),f'{i+1:02d} / 13',font=font,fill='#27616b')
canvas.save(ROOT/'.runtime/ppt-contact-sheet.jpg')
(ROOT/'docs/evidence/documents-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8')
print('Verified Word/PDF pairs:',len(report['documents']),'and',len(prs.slides),'slides; contact sheet generated.')
