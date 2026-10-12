"""Generate editable Word reports and a presentation from the actual project evidence."""
from pathlib import Path
import sys,re,json,zipfile
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.tools'))
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PI,Pt as PP
from pptx.dml.color import RGBColor as PC
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
OUT=ROOT/'deliverables';OUT.mkdir(exist_ok=True)
TEAL='087F83';DARK='174950';MUTED='68868D';BG='F4F8F8'
summary=json.loads((ROOT/'docs/evidence/summary.json').read_text('utf-8'))

def plain(s):
    s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1（\2）',s)
    return s.replace('**','').replace('`','').replace('<br>','\n')

def font_run(run,size=None,bold=None,color=None):
    run.font.name='Microsoft YaHei'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    if color:run.font.color.rgb=RGBColor.from_string(color)

def word_base(title,subtitle):
    doc=Document();sec=doc.sections[0]
    sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
        style=doc.styles[name];style.font.name='Microsoft YaHei'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    doc.styles['Normal'].font.size=Pt(10.5)
    doc.styles['Normal'].paragraph_format.line_spacing=1.2
    for name,size in [('Heading 1',19),('Heading 2',14),('Heading 3',12)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].font.color.rgb=RGBColor.from_string(TEAL)
    h=sec.header.paragraphs[0];font_run(h.add_run('鲜境 FreshMarket  /  AI辅助软件开发实践 · 实验四'),9,color=MUTED)
    footer=sec.footer.paragraphs[0];footer.alignment=2
    font_run(footer.add_run('FreshMarket · 2026-10-06   |   '),8,color=MUTED)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.add_paragraph('FRESHMARKET  /  PROJECT DELIVERY','Subtitle')
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(65);font_run(p.add_run(title),28,True,DARK)
    doc.add_paragraph(subtitle,'Subtitle')
    doc.add_paragraph('独立品牌 · 仿盒马生鲜选购体验 · 实际持久化业务闭环')
    doc.add_paragraph('版本 1.0  |  编制日期 2026-10-06')
    doc.add_paragraph('依据实验任务书与教材第九章；记录本次实际实现及测试范围。')
    doc.add_paragraph('提交人姓名 / 学号 / 专业：________________________')
    doc.add_page_break();return doc

def table(doc,rows):
    if not rows:return
    count=max(map(len,rows));t=doc.add_table(rows=0,cols=count);t.style='Light Shading Accent 1'
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,cell in enumerate(cells):
            cell.text=plain(row[j] if j<len(row) else '')
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4)
                for r in p.runs:font_run(r,8.5 if count>5 else 9.5,i==0)
        if i==0:
            prop=t.rows[0]._tr.get_or_add_trPr();prop.append(OxmlElement('w:tblHeader'))
        else:
            prop=t.rows[i]._tr.get_or_add_trPr();prop.append(OxmlElement('w:cantSplit'))
    doc.add_paragraph()

def markdown(doc,text,base):
    lines=text.splitlines();i=0
    while i<len(lines):
        s=lines[i].strip();i+=1
        if not s:continue
        if s.startswith('```'):
            language=s[3:];block=[]
            while i<len(lines) and not lines[i].strip().startswith('```'):block.append(lines[i]);i+=1
            i+=1
            if language!='mermaid':
                p=doc.add_paragraph();r=p.add_run('\n'.join(block));r.font.name='Consolas';r.font.size=Pt(8.5)
            continue
        if s.startswith('|'):
            rows=[]
            while True:
                row=[p.strip() for p in s.strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',p) for p in row):rows.append(row)
                if i>=len(lines) or not lines[i].strip().startswith('|'):break
                s=lines[i].strip();i+=1
            table(doc,rows);continue
        img=re.match(r'!\[([^\]]*)\]\(([^)]+)\)',s)
        if img:
            file=base/img[2];w,h=Image.open(file).size;scale=min(6.65/w,7.1/h)
            doc.add_picture(str(file),width=Inches(w*scale),height=Inches(h*scale))
            p=doc.add_paragraph(img[1]+' · 完整原图：'+img[2]);p.alignment=1
            for r in p.runs:font_run(r,8,color=MUTED)
            continue
        heading=re.match(r'(#{1,4})\s+(.+)',s)
        if heading:
            if len(heading[1])==1:continue
            doc.add_heading(plain(heading[2]),min(3,len(heading[1])-1));continue
        if s.startswith('- '):doc.add_paragraph(plain(s[2:]),'List Bullet');continue
        if re.match(r'\d+\. ',s):doc.add_paragraph(plain(re.sub(r'^\d+\. ','',s)),'List Number');continue
        p=doc.add_paragraph(plain(s.lstrip('> ')))
        if s.startswith('>'):p.paragraph_format.left_indent=Inches(.2)

lab=f'''# AI辅助软件开发实验报告（四）

## 项目与提交信息
项目题目：【项目实战】企业级应用——鲜境生鲜商城系统。
姓名：________________   学号：________________   专业：________________
提交日期：________________   小组成员与实际分工：________________
编程工具：Codex、Node.js 24、npm workspaces、TypeScript、Docker Desktop、Playwright。
大模型：GPT-6系列（本会话系统标识；具体服务端构建号未公开）。

## 一、实验目的
根据教材第九章社区超市商城思路，完成企业级商城从需求、选型、数据建模、编码、测试到部署与工程化交付的完整实践。以生鲜购物为业务主题，验证购物闭环、角色职责、库存及资金沙箱一致性，并留下可复现证据。

## 二、实验步骤展示
### 2.1 立项与需求分析
系统名为“鲜境 FreshMarket”，面向城市家庭、通勤顾客和生鲜运营团队。采用独立品牌，借鉴盒马的分类选购、规格产地展示、冷链与到家体验。6类24款食材；游客及6种账号角色；完整用例、流程、领域模型、ER、权限矩阵、API与可验证指标见需求文档。
![生鲜购物界面](screenshots/01-home-desktop.png)

### 2.2 技术选型与环境搭建
React与Fastify为独立前后端工程，TypeScript统一开发。Docker实际启动Nginx、API、PostgreSQL、Redis、RabbitMQ、Prometheus、Grafana共7个容器。核心业务模块分层，配置中心由环境变量承担；模块化单体规模无需服务注册发现。精确依赖见锁文件，替代方案与风险见设计文档、ADR。
![实际部署环境](screenshots/07-runtime-stack.png)
![实际分层架构](diagrams/architecture.png)

### 2.3 AI编程过程与关键决策
原始Prompt：按照所给文档的要求完成系统搭建，目标是仿盒马搭建企业级生鲜购物平台。
工具/模型、阶段输入、生成内容、采纳理由、测试结果与实际修正记录见AI过程记录。采用整数分金额、服务端价格、下单快照、事务条件扣库存、行锁和唯一约束。采用Outbox同事务事件及RabbitMQ确认/去重，记录认证、缓存、消息及部署决策。个人与团队人工修改、评审和分工需由提交人补录。

### 2.4 业务实现与企业级能力
浏览→搜索→详情→购物车→地址/配送时段→幂等下单→沙箱付款→仓库拣货/配送→顾客收货/评价→财务结算。待付款取消及15分钟过期释放库存；整单退款审核冻结履约和结算。已配送生鲜不回库存。接口JWT、密码scrypt哈希、RBAC、所有权校验、参数化SQL、输入校验、限流、审计及traceId；OpenAPI、Docker、监控、CI配置、备份/回滚脚本均已交付。
![真实越权拦截](screenshots/10-rbac-denied.png)
![实际监控](screenshots/08-grafana-monitor.png)

### 2.5 测试验证
| 验证 | 最终实际结果 |
| --- | --- |
| SQLite / PostgreSQL | 各28/28通过 |
| SQLite覆盖率 | 行{summary['sqliteCoverage']['line']}%，分支{summary['sqliteCoverage']['branch']}% |
| PostgreSQL覆盖率 | 行{summary['postgresCoverage']['line']}%，分支{summary['postgresCoverage']['branch']}% |
| 浏览器E2E | 本地5/5，完整Docker5/5 |
| 实际Redis/RabbitMQ链路 | 8/8通过 |
| 依赖审计 | 0项已知告警 |
| 类型检查 / 生产构建 | 通过 |
| 本机商品接口短时测量 | 200请求，并发10，p95 {summary['load']['p95Ms']}ms，错误率0% |
最后一项总时长约0.168秒，是本机短时突发测量；不代表持续容量或生产SLA。CI提供配置但未声明GitHub实际执行；恢复脚本未破坏性演练。

### 2.6 交付与演示
源码、SQL、部署脚本、Word需求/设计/测试报告、实验报告、PPT、实际浏览器WebM录屏、原始测试证据及SHA256均交付。视频含购物闭环、真实403、沙箱重复支付/退款、架构/Outbox说明、监控、测试、AI使用与复现；录屏字幕是讲解层。个人信息、真人讲解和手写签名由提交人完成。

## 三、开发亮点、问题及解决方案
| 难点 | 实际处理与证据 |
| --- | --- |
| 并发库存 / 订单幂等 | 锁用户与订单；条件扣减；8次幂等提交、最后3件争抢测试 |
| 付款、取消、退款和结算竞争 | 状态机、行锁与唯一约束；重复操作和互斥测试 |
| 业务提交与消息投递可靠性 | 同事务Outbox；ConfirmChannel；持久化消费去重；8项实物组件验证 |
| 扫描教材无法直接读取 | 渲染指定章节页面，视觉读取实际PDF150–156页 |
| Windows子进程 / Docker管道权限 | 使用执行工具权限流程；实际启动Docker并继续验证 |
| E2E登录定位歧义 | 限定header按钮，保留失败原因并回归 |
| Swagger关联依赖告警 | 升级Swagger UI 6.1.1；重新审计及双数据库回归 |

## 四、小结
### 4.1 技术收获与体会
本项目的实践结论是：商城设计需围绕业务不变量建立数据约束，金额与库存由服务端负责；前端菜单权限必须与接口权限结合；事务内远程消息发布存在可靠性缺口，应采用持久化Outbox与消费去重；功能截图应与真实测试、部署和日志共同证明实现。AI生成代码需经类型检查、并发与异常测试、依赖审计和真实组件验证后采纳。
个人或团队学习体会（提交人补充）：________________________________________

### 4.2 今后改进方向
接入真实支付签名回调与退款渠道、实际配送和地理围栏、批次溯源与保质期管理；完善部分退款、已结算逆向账务、优惠券及多门店；验收集群限流、HTTPS、故障注入、长时压测、数据库高可用与备份恢复演练。当前实验功能边界在需求/设计/测试报告均有注明。

## 小组成员手写电子签名
成员姓名：________________  实际分工：________________  手写签名：________________
成员姓名：________________  实际分工：________________  手写签名：________________
以上信息须由小组据实填写，不由生成工具代签。
'''
(ROOT/'docs/lab-report.md').write_text(lab,'utf-8')
reports=[('requirements','需求文档','用户画像、权限与可验证业务需求'),('design','设计文档','模块化架构与关键一致性设计'),('test-report','测试报告','双数据库与实际组件验证证据'),('lab-report','实验报告','实验四：AI辅助企业级生鲜商城实践'),('deployment','部署运维手册','配置分离、启动、监控、备份与回滚'),('ai-log','AI过程与决策记录','实际Prompt、采纳依据与修正记录')]
for filename,title,subtitle in reports:
    doc=word_base('鲜境 · '+title,subtitle)
    markdown(doc,(ROOT/f'docs/{filename}.md').read_text('utf-8'),ROOT/'docs')
    if filename=='ai-log':doc.add_page_break();doc.add_heading('架构决策记录 ADR',1);markdown(doc,(ROOT/'docs/adr.md').read_text('utf-8'),ROOT/'docs')
    doc.core_properties.title='鲜境 FreshMarket '+title;doc.core_properties.author='FreshMarket 项目交付'
    doc.save(OUT/f'鲜境_{title}.docx')

prs=Presentation();prs.slide_width=PI(13.333);prs.slide_height=PI(7.5)

def txt(slide,x,y,w,h,content,size=20,color=DARK,bold=False):
    tf=slide.shapes.add_textbox(PI(x),PI(y),PI(w),PI(h)).text_frame;tf.word_wrap=True
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.space_after=PP(12)
        p.font.name='Microsoft YaHei';p.font.size=PP(size);p.font.bold=bold;p.font.color.rgb=PC.from_string(color)
    return tf

def box(slide,x,y,w,h,color):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PI(x),PI(y),PI(w),PI(h));s.adjustments[0]=.07;s.fill.solid();s.fill.fore_color.rgb=PC.from_string(color);s.line.fill.background();return s

def slide(title,subtitle,notes):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=PC.from_string(BG)
    box(s,0,0,13.333,.12,TEAL);txt(s,.55,.33,12,.35,'FRESHMARKET  /  ENTERPRISE FRESH GROCERY',11,TEAL,True)
    txt(s,.55,.9,12, .7,title,29,DARK,True);txt(s,.57,1.6,12,.45,subtitle,13,MUTED)
    txt(s,.55,7.12,11,.25,'鲜境 · 实验四 · 2026-10-06  |  实际系统与证据，支付/配送为实验沙箱',9,MUTED)
    txt(s,12.15,7.1,.65,.25,f'{len(prs.slides):02d}',10,TEAL,True)
    s.notes_slide.notes_text_frame.text=notes;return s

def pic(s,path,x,y,w,h):
    p=ROOT/path;iw,ih=Image.open(p).size;factor=min(w/iw,h/ih);pw,ph=iw*factor,ih*factor
    if path.endswith('01-home-desktop.png'):
        shape=s.shapes.add_picture(str(p),PI(x),PI(y),width=PI(w),height=PI(h))
        shape.crop_bottom=max(0,1-(h/w)/(ih/iw));return
    s.shapes.add_picture(str(p),PI(x+(w-pw)/2),PI(y+(h-ph)/2),width=PI(pw),height=PI(ph))

def visual(title,subtitle,path,bullets,notes):
    s=slide(title,subtitle,notes);pic(s,path,.55,2.2,8.4,4.6)
    box(s,9.18,2.25,3.6,4.48,'FFFFFF');txt(s,9.45,2.55,3.05,3.8,bullets,17);return s

s=slide('鲜境 FreshMarket','仿盒马生鲜购物体验 · 企业级商城工程实践','开场说明：独立鲜境品牌。讲解对象为实际完成的实验平台，资金和配送使用沙箱。姓名学号由提交人填写。')
txt(s,.7,2.5,7,1.7,'把新鲜带回家，\n把业务闭环落到系统。',33,DARK,True)
txt(s,.7,4.6,7,1.5,'6类食材 · 24款商品 · 6种账号角色\n商城 / 工作台 / 数据库 / 消息 / 监控\n姓名、学号、小组成员：________________',18,MUTED)
pic(s,'web/public/images/hero.jpg',8.4,2.25,4.1,4.45)
visual('01 目标用户与角色职责','城市家庭、通勤鲜客与生鲜门店运营团队','docs/diagrams/use-cases.png','游客：发现好食材\n会员：购物与售后\n运营 / 仓库 / 客服 / 财务\n管理员：授权与审计','说明6种账号和游客权限，个人购物资源始终要求本人。菜单隐藏不替代API鉴权。')
visual('02 顾客选购体验','分类、搜索、规格产地与冷链说明；桌面及手机可用','docs/screenshots/01-home-desktop.png','6类 / 24款食材\n本地图片与来源记录\n搜索 / 排序 / 分页\n390px手机E2E通过','打开实际商城演示搜索三文鱼、详情、购物车。示意图片和商品数据为实验种子，前端通过真实JSON API查询。')
visual('03 完整购物业务闭环','每一步均有数据状态、所有权与角色校验','docs/diagrams/business-flow.png','下单预扣库存\n15分钟付款期限\n拣货 / 配送 / 收货\n评价 / 售后 / 财务','在录屏中展示实际下单与付款，切换仓库和会员完成履约；第二单演示整单退款。已配送的生鲜不能自动回可售库存。')
visual('04 权限与审计','实际仓库账号调用用户管理接口，服务端返回403','docs/screenshots/10-rbac-denied.png','JWT主体 + DB角色\nscrypt随机盐密码\n本人订单 / 地址校验\ntraceId贯通审计','解释实际403证据。修改角色会影响旧JWT。角色注册字段拒绝注入，SQL参数绑定，接口限流。公网HTTPS需正式部署配置。')
visual('05 部署与模块分层','前端JSON交互；Controller → Service → Repository','docs/diagrams/architecture.png','React / Fastify / TS\nPostgreSQL事务\nRedis查询缓存\nRabbitMQ可靠事件','采用模块化单体来降低当前规模复杂度。SQLite本地可立即启动，完整Docker采用PG；配置中心为环境变量，不部署无需使用的注册发现服务。')
visual('06 数据模型与不变量','15张表；金额整数分；订单价格和地址快照','docs/diagrams/er.png','付款 / 退款 / 结算\n订单id唯一约束\n用户 + key唯一\n索引与连接池','完整DDL与初始化SQL随源码交付。Outbox和processed_events为逻辑事件对应，不设置物理FK；订单地址为JSON快照。')
s=slide('07 库存与状态竞争的处理','行锁 + 条件更新 + 唯一约束 + 同事务审计/事件','重点讲解库存争抢、重复下单和付款取消竞争测试。不能用Redis缓存值做最终库存决定。')
box(s,.6,2.35,7.7,2.15,'174950');txt(s,.88,2.58,7.15,1.65,'UPDATE products\nSET stock = stock - quantity\nWHERE stock >= quantity AND active = 1',20,'FFFFFF')
txt(s,.7,4.87,7.5,1.55,'按商品id排序扣减，任何一项失败整单回滚。\n先锁用户串行购物车与下单；再检查幂等摘要。\n付款 / 取消锁定同一订单，只有一个有效转换。',18)
txt(s,8.95,2.4,3.6,4.3,'8次并发提交 → 1单\n最后3件库存争抢 → 1单\n6次重复支付 → 1笔\n售后pending → 冻结履约/结算',20,TEAL,True)
s=slide('08 Outbox与实际消息链路','处理业务数据库与RabbitMQ之间的可靠性缺口','8项实物组件验证已执行。confirm确认后标记sent；消费者事务提交后ack；重复事件不会重复产生通知。')
for x,num,title,body in [( .7,'01','业务事务','订单 / 审计 / Outbox\n一次事务提交'),(3.9,'02','确认发布','processing租约\n持久化 + Confirm'),(7.1,'03','去重消费','processed_events唯一\n通知提交后Ack'),(10.3,'04','失败恢复','重试 / 重连\n死信队列检查')]:
    box(s,x,2.65,2.5,3.15,'FFFFFF');txt(s,x+.22,2.92,2,.5,num,27,TEAL,True);txt(s,x+.22,3.67,2,.55,title,20,DARK,True);txt(s,x+.22,4.4,2.1,1.15,body,16,MUTED)
txt(s,.75,6.2,12,.5,'实际检查：缓存TTL / 共享失效 / PG下单 / 发布确认 / 消费落库 / 重复去重 / 取消回库 / 消息投递',14)
s=slide('09 测试与实际验证结果','保留双数据库日志、E2E报告、依赖审计及测量原始数据','用例包含并发、异常、越权、回滚、超时、退款、结算和权限撤销。200次短时测量不作为生产SLA。')
for x,y,title,value,detail in [(.7,2.4,'SQLite / PostgreSQL','28 / 28','两个环境各28项通过'),(4.85,2.4,'浏览器 E2E','5 / 5','本地、Docker分别通过'),(9,2.4,'实物组件','8 / 8','真实Redis / RabbitMQ'),(.7,4.55,'行覆盖率','87.02%','SQLite；PG 88.20%'),(4.85,4.55,'依赖审计','0 告警','修复Swagger关联依赖'),(9,4.55,'本机短测 p95','17.34 ms','200请求 / 并发10 / 0错误')]:
    box(s,x,y,3.65,1.85,'FFFFFF');txt(s,x+.25,y+.16,3.1,.4,title,15,MUTED);txt(s,x+.25,y+.61,3.1,.65,value,29,TEAL,True);txt(s,x+.25,y+1.36,3.15,.32,detail,12,MUTED)
visual('10 部署、可观测性与运维','七个容器实际运行；Prometheus target UP = 1','docs/screenshots/08-grafana-monitor.png','实际QPS / p95 / 内存\n健康、日志与告警规则\n配置分离 / 镜像回滚\n备份与恢复脚本','展示真实Grafana及健康接口。CI配置已交付未在GitHub执行，恢复脚本未破坏性演练。课堂Compose仅本机HTTP，正式公网需TLS与真实渠道。')
s=slide('11 AI辅助开发与修正记录','记录原始Prompt、工具/模型、生成内容、采纳理由和实际验证','Codex / GPT-6系列。用户原始Prompt是本次真实输入；其他阶段任务为代理归纳，未伪造用户对话、团队评审和签名。')
box(s,.65,2.3,12,1,'E3F2EF');txt(s,.9,2.54,11.45,.6,'“按照所给文档的要求完成系统搭建，目标是仿盒马搭建企业级生鲜购物平台”',19,TEAL,True)
txt(s,.8,3.75,5.65,2.7,'采纳：TS模块化单体、整数分金额\n采纳：快照 / 行锁 / Outbox去重\n证据：AI过程表与ADR\n源码：锁文件 + SQL + Compose',19)
txt(s,7,3.75,5.35,2.7,'扫描PDF → 渲染章节视觉读取\nWindows权限 → 执行工具审批\nE2E定位歧义 → 精确限定header\nSwagger告警 → 升级并重新回归',19)
s=slide('12 交付、复现与后续方向','可运行源码 + 三类Word文档 + 实验报告 + PPT + 实际系统录屏','结束说明当前已经启动。交付人补个人信息和手写电子签名；真实资金、配送、HA、长时容量及恢复演练是后续工作。')
box(s,.65,2.3,6.3,3.6,'174950');txt(s,.94,2.61,5.7,2.9,'本地：Node.js 24+\nnpm ci → npm run dev\n\nDocker：生成 .env.docker\ndocker compose --env-file .env.docker\nup -d --build',18,'FFFFFF')
txt(s,7.5,2.45,5.1,3.8,'已交付：Word / PPT / WebM / 源码ZIP\n完整Docker商城：localhost:8080\n后续：真实渠道、部分退款、批次溯源\n验收：集群限流、HA、恢复与长时压测\n提交人：补姓名、学号、分工及签名',19)
prs.core_properties.title='鲜境 FreshMarket 企业级生鲜商城答辩';prs.core_properties.author='FreshMarket 项目交付'
deck=OUT/'鲜境_项目汇报.pptx';prs.save(deck)
manifest={'wordDocuments':[f'鲜境_{t}.docx' for _,t,_ in reports],'presentation':deck.name,'slides':len(prs.slides),'generatedFrom':'docs/*.md, docs/evidence/summary.json, actual screenshots','date':'2026-10-06'}
for file in list(OUT.glob('*.docx'))+[deck]:
    with zipfile.ZipFile(file) as z:assert z.testzip() is None
(OUT/'documents-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),'utf-8')
print(f'Generated {len(reports)} editable Word reports and {len(prs.slides)} presentation slides; OOXML archives verified.')
